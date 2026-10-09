"""Find KOLs → Export: the column catalogue and the CSV / Excel writers.

One catalogue serves both the directory and the pending-verification list; a
column with no value for a row (e.g. "X Handle" on a Practo KOL) is left blank.
"Status" reads "Approved" for directory KOLs.
"""
import csv
import io
from datetime import datetime
from typing import Callable, Dict, List, Tuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SOURCE_LABELS = {"manual": "Manual", "import": "Imported", "hcp": "HCP profile",
                 "twitter": "X (Twitter)", "practo": "Practo", "google": "Google search"}


def _name(r) -> str:
    return " ".join(p for p in (getattr(r, "title", None), r.first_name, r.last_name) if p)


def _enrich(r, block: str):
    data = r.enrichment or {}
    return ((data.get(block) or {}).get("total_count")) if isinstance(data, dict) else None


def _date(v) -> str:
    return v.strftime("%Y-%m-%d %H:%M") if isinstance(v, datetime) else ""


# key -> (header, getter, on by default)
COLUMNS: Dict[str, Tuple[str, Callable, bool]] = {
    "name": ("Name", _name, True),
    "title": ("Title", lambda r: getattr(r, "title", None), False),
    "first_name": ("First Name", lambda r: r.first_name, False),
    "last_name": ("Last Name", lambda r: r.last_name, False),
    "email": ("Email", lambda r: r.email, True),
    "contact_no": ("Contact No.", lambda r: r.contact_no, True),
    "therapeutic_area": ("Therapeutic Area", lambda r: r.therapeutic_area, True),
    "specialty": ("Specialty", lambda r: r.specialty, True),
    "subspecialty": ("Subspecialty", lambda r: r.subspecialty, False),
    "designation": ("Designation", lambda r: r.designation, True),
    "institution": ("Society / Hospital / Institution", lambda r: r.institution, True),
    "experience_years": ("Total Experience (years)", lambda r: r.experience_years, True),
    "influence_level": ("Influence Level", lambda r: r.influence_level, True),
    "functional_expertise": ("Functional Expertise", lambda r: r.functional_expertise, False),
    "geographic_zone": ("Country", lambda r: r.geographic_zone, True),
    "state_city": ("State / City", lambda r: r.state_city, True),
    "publications": ("Publications (PubMed)", lambda r: _enrich(r, "pubmed"), False),
    "clinical_trials": ("Clinical Trials", lambda r: _enrich(r, "clinical_trials"), False),
    "source": ("Source", lambda r: SOURCE_LABELS.get(r.source, r.source), True),
    "status": ("Status", lambda r: (getattr(r, "status", None) or "approved").capitalize(), False),
    "x_username": ("X Handle", lambda r: f"@{r.x_username}" if r.x_username else None, False),
    "practo_profile_url": ("Practo Profile", lambda r: r.practo_profile_url, False),
    "web_profile_url": ("Web Page", lambda r: r.web_profile_url, False),
    "notes": ("Notes", lambda r: r.notes, False),
    "created_at": ("Added On", lambda r: _date(r.created_at), False),
}


def catalogue() -> List[Dict]:
    return [{"key": k, "label": label, "default": default} for k, (label, _, default) in COLUMNS.items()]


def _rows(records, keys: List[str]):
    getters = [COLUMNS[k][1] for k in keys]
    for r in records:
        yield ["" if (v := g(r)) is None else v for g in getters]


def to_csv(records, keys: List[str]) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([COLUMNS[k][0] for k in keys])
    w.writerows(_rows(records, keys))
    # BOM so Excel opens UTF-8 (names with non-ASCII characters) correctly.
    return ("﻿" + buf.getvalue()).encode("utf-8")


def to_xlsx(records, keys: List[str], title: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append([COLUMNS[k][0] for k in keys])
    widths = [len(COLUMNS[k][0]) for k in keys]
    for row in _rows(records, keys):
        ws.append(row)
        for i, v in enumerate(row):
            widths[i] = max(widths[i], min(60, len(str(v))))
    head_fill = PatternFill("solid", fgColor="EEF4E0")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = head_fill
        cell.alignment = Alignment(vertical="center")
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w + 2
    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
